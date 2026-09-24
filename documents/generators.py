"""Генерация документов DOCX и XLSX."""
from io import BytesIO

from django.conf import settings
from django.utils import timezone
from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from documents.models import GeneratedDocument
from hospital.models import Feedback, Schedule


class DocumentGenerator:
    """Обёртка для генерации документов."""

    @staticmethod
    def _save_record(user, doc_type, title, parameters=None):
        return GeneratedDocument.objects.create(
            user=user,
            document_type=doc_type,
            title=title,
            parameters=parameters or {},
        )

    @classmethod
    def generate_doctor_schedule_docx(cls, doctor, user):
        """Выписка из расписания врача в формате DOCX."""
        document = Document()
        title = document.add_heading(settings.SITE_NAME, level=0)
        title.alignment = 1

        document.add_heading(f'Расписание приёма: {doctor.full_name}', level=1)
        document.add_paragraph(f'Специальность: {doctor.specialty}')
        document.add_paragraph(f'Отделение: {doctor.department.name}')
        document.add_paragraph(f'Кабинет: {doctor.cabinet or "—"}')
        document.add_paragraph(f'Дата формирования: {timezone.now().strftime("%d.%m.%Y %H:%M")}')

        table = document.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        headers = ['День недели', 'Начало', 'Конец', 'Кабинет']
        for index, header_text in enumerate(headers):
            header_cells[index].text = header_text
            for paragraph in header_cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        schedules = Schedule.objects.filter(doctor=doctor).order_by('day_of_week', 'start_time')
        for schedule in schedules:
            row_cells = table.add_row().cells
            row_cells[0].text = schedule.get_day_of_week_display()
            row_cells[1].text = schedule.start_time.strftime('%H:%M')
            row_cells[2].text = schedule.end_time.strftime('%H:%M')
            row_cells[3].text = schedule.room or doctor.cabinet or '—'

        if not schedules.exists():
            document.add_paragraph('Расписание пока не заполнено.')

        footer = document.add_paragraph()
        footer.add_run('\n\nДокумент сформирован автоматически системой ').font.size = Pt(10)
        footer_run = footer.add_run(settings.SITE_NAME)
        footer_run.font.size = Pt(10)
        footer_run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)

        record_title = f'Расписание {doctor.full_name}'
        cls._save_record(user, GeneratedDocument.TYPE_SCHEDULE, record_title, {'doctor_id': doctor.pk})

        return buffer, f'raspisanie_{doctor.slug}.docx'

    @classmethod
    def generate_feedback_report_xlsx(cls, user, date_from=None, date_to=None):
        """Отчёт по заявкам обратной связи за период."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Заявки'

        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='0066CC', end_color='0066CC', fill_type='solid')
        headers = ['ID', 'Дата', 'Имя', 'Email', 'Телефон', 'Тема', 'Сообщение', 'Статус']

        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        queryset = Feedback.objects.all().order_by('-created_at')
        if date_from:
            queryset = queryset.filter(created_at__date__gte=date_from)
        if date_to:
            queryset = queryset.filter(created_at__date__lte=date_to)

        status_map = dict(Feedback.STATUS_CHOICES)
        for row_num, feedback in enumerate(queryset, start=2):
            sheet.cell(row=row_num, column=1, value=feedback.pk)
            sheet.cell(row=row_num, column=2, value=feedback.created_at.strftime('%d.%m.%Y %H:%M'))
            sheet.cell(row=row_num, column=3, value=feedback.name)
            sheet.cell(row=row_num, column=4, value=feedback.email)
            sheet.cell(row=row_num, column=5, value=feedback.phone)
            sheet.cell(row=row_num, column=6, value=feedback.subject)
            sheet.cell(row=row_num, column=7, value=feedback.message[:500])
            sheet.cell(row=row_num, column=8, value=status_map.get(feedback.status, feedback.status))

        for column in sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            sheet.column_dimensions[column_letter].width = min(max_length + 2, 50)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        period = f'{date_from or "all"}_{date_to or "all"}'
        record_title = f'Отчёт по заявкам ({period})'
        cls._save_record(
            user,
            GeneratedDocument.TYPE_FEEDBACK_REPORT,
            record_title,
            {'date_from': str(date_from), 'date_to': str(date_to)},
        )

        filename = f'otchet_zayavki_{timezone.now().strftime("%Y%m%d")}.xlsx'
        return buffer, filename
